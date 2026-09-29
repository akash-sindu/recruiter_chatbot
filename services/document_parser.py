from pypdf import PdfReader
from docx import Document


def extract_text_from_file(file_storage) -> str:
    """Extracts raw text from PDF, DOCX, or TXT uploads."""
    filename = file_storage.filename.lower()

    if filename.endswith(".pdf"):
        reader = PdfReader(file_storage)
        text = ""
        for page in reader.pages:
            page_text = page.extract_text()
            if page_text:
                text += page_text + "\n"
        return text

    elif filename.endswith(".docx"):
        doc = Document(file_storage)
        text = ""
        for para in doc.paragraphs:
            if para.text.strip():
                text += para.text + "\n"
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        if cell.text.strip():
                            text += cell.text + "\n"
        return text
    elif filename.endswith(".txt"):
        return file_storage.read().decode("utf-8")

    else:
        raise ValueError("Unsupported format. Upload PDF, DOCX, or TXT.")
